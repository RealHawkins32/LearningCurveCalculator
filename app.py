"""SEC EDGAR Learning-Curve Calculator — Streamlit application.

Financial figures are extracted from SEC XBRL companyfacts, not inferred.
Unit-volume inputs must be supplied and documented by the user.
"""
import io
import math
import os
import re
from datetime import date

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import plotly.graph_objects as go
import requests
import streamlit as st
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
from docx import Document
from docx.shared import Inches, Pt, RGBColor

st.set_page_config(page_title="Learning-Curve Calculator", page_icon="📈", layout="wide")
st.title("📈 Learning-Curve Calculator")
st.caption("SEC EDGAR financial data • Documented unit estimates • Experience-curve regression • Downloadable Word report")

# SEC requires an identifiable user agent. Set SEC_CONTACT_EMAIL in .streamlit/secrets.toml
# or as an environment variable before deploying publicly.
def setting(name, default=""):
    try:
        return str(st.secrets.get(name, os.getenv(name, default)))
    except Exception:
        return os.getenv(name, default)

sec_email = setting("SEC_CONTACT_EMAIL")
if not sec_email or "@" not in sec_email:
    st.error("Configure SEC_CONTACT_EMAIL with your contact email in .streamlit/secrets.toml before using the app.")
    st.code('SEC_CONTACT_EMAIL = "your.email@example.com"')
    st.stop()

HEADERS = {
    "User-Agent": f"LearningCurveCalculator/1.0 ({sec_email})",
    "Accept-Encoding": "gzip, deflate",
}

@st.cache_data(ttl=86400, show_spinner=False)
def get_tickers():
    r = requests.get("https://www.sec.gov/files/company_tickers.json", headers=HEADERS, timeout=30)
    r.raise_for_status()
    return r.json()

@st.cache_data(ttl=86400, show_spinner=False)
def get_facts(cik):
    r = requests.get(f"https://data.sec.gov/api/xbrl/companyfacts/CIK{cik:010d}.json", headers=HEADERS, timeout=45)
    r.raise_for_status()
    return r.json()

def find_company(ticker):
    return next((x for x in get_tickers().values() if x["ticker"].upper() == ticker.upper()), None)

COGS_TAGS = ["CostOfGoodsAndServicesSold", "CostOfRevenue", "CostOfGoodsSold"]
REV_TAGS = ["RevenueFromContractWithCustomerExcludingAssessedTax", "Revenues", "SalesRevenueNet", "RevenueFromContractWithCustomerIncludingAssessedTax"]

def annual_facts(facts, tags, count):
    """Choose earliest 10-K annual fact for each period, preferring higher-priority tag.

    NOTE: earliest *available* filing may be a later comparative 10-K.
    """
    gaap = facts.get("facts", {}).get("us-gaap", {})
    for tag in tags:
        if tag not in gaap:
            continue
        rows = []
        for rec in gaap[tag].get("units", {}).get("USD", []):
            if rec.get("form") != "10-K" or not rec.get("start") or not rec.get("end"):
                continue
            duration = (pd.Timestamp(rec["end"]) - pd.Timestamp(rec["start"])).days
            if not 330 <= duration <= 380:
                continue
            if rec.get("val") is None or rec["val"] <= 0:
                continue
            rows.append({"Fiscal Year": pd.Timestamp(rec["end"]).year,
                         "Period End": rec["end"], "Value": float(rec["val"]),
                         "Filing Date": rec.get("filed", ""), "Accession": rec.get("accn", ""),
                         "XBRL Tag": tag})
        if rows:
            df = pd.DataFrame(rows).sort_values(["Period End", "Filing Date"])
            df = df.drop_duplicates("Period End", keep="first").sort_values("Period End")
            return df.tail(count).reset_index(drop=True)
    return pd.DataFrame()

def accession_link(cik, accession):
    if not accession:
        return ""
    return f"https://www.sec.gov/Archives/edgar/data/{cik}/{accession.replace('-', '')}/"

def regression(frame):
    x = np.log(frame["Cumulative Units"].to_numpy(dtype=float))
    y = np.log(frame["Unit Cost ($)"].to_numpy(dtype=float))
    if len(x) < 5 or np.ptp(x) < 1e-12:
        raise ValueError("At least five observations with varying cumulative units are required.")
    slope, intercept = np.polyfit(x, y, 1)
    fitted = intercept + slope * x
    ss_tot = np.sum((y - np.mean(y)) ** 2)
    r2 = 1 - np.sum((y - fitted) ** 2) / ss_tot if ss_tot > 1e-15 else 0.0
    b = -float(slope)
    learning = 1 - 2 ** (-b)
    return b, float(learning), float(r2), np.exp(fitted), float(np.exp(intercept))

def assessment(b, learning, r2, method):
    if r2 >= 0.8:
        fit = "The log-log regression fits the observed data strongly (R² ≥ 0.80)."
    elif r2 < 0.5:
        fit = "The fit is weak (R² < 0.50); no reliable experience curve is visible in this window."
    else:
        fit = "The fit is mixed (0.50 ≤ R² < 0.80), so the inferred learning rate is uncertain."
    if learning > 0:
        direction = f"The estimated learning rate is {learning * 100:.1f}% lower unit cost per doubling of measured cumulative units."
    elif learning < 0:
        direction = f"The estimated learning rate is {learning * 100:.1f}%, meaning unit cost increased rather than decreased as measured cumulative units doubled."
    else:
        direction = "The estimated learning rate is approximately zero."
    caution = ("Unit counts are estimates derived from revenue and assumed annual average selling prices; uncertain prices and product mix can distort the apparent curve."
               if method == "Revenue / estimated annual average selling prices" else
               "Unit counts were supplied by the user and require verification against the cited disclosures.")
    return (f"{direction} {fit} This pattern is descriptive and does not by itself establish that accumulated production caused the cost change. "
            f"{caution} Cumulative units start at the beginning of the selected data window, not the company's lifetime production; the results alone cannot prove a durable scale advantage or industry-leading learning rate.")

def plotly_chart(frame, fitted):
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=frame["Cumulative Units"], y=frame["Unit Cost ($)"], mode="markers+lines", name="Observed unit cost"))
    fig.add_trace(go.Scatter(x=frame["Cumulative Units"], y=fitted, mode="lines", name="Fitted power curve", line=dict(dash="dash")))
    fig.update_layout(title="Experience curve (log-log)", xaxis_title="Cumulative units (log scale)", yaxis_title="Unit cost ($, log scale)", height=440)
    fig.update_xaxes(type="log")
    fig.update_yaxes(type="log")
    return fig

def png_chart(frame, fitted):
    fig, ax = plt.subplots(figsize=(8.5, 4.2))
    ax.plot(frame["Cumulative Units"], frame["Unit Cost ($)"], "o-", label="Observed")
    ax.plot(frame["Cumulative Units"], fitted, "--", label="Fitted power curve")
    ax.set(xscale="log", yscale="log", xlabel="Cumulative units (log)", ylabel="Unit cost ($, log)", title="Experience curve: observed vs fitted")
    ax.grid(True, which="both", alpha=0.25)
    ax.legend()
    fig.tight_layout()
    out = io.BytesIO()
    fig.savefig(out, format="png", dpi=155)
    plt.close(fig)
    out.seek(0)
    return out

def gemini_analysis(company, ticker, frame, b, learning, r2, method, base_text):
    api_key = setting("GEMINI_API_KEY")
    if not api_key:
        return None, "No GEMINI_API_KEY configured; the verified rule-based interpretation remains available."
    model = setting("GEMINI_MODEL", "gemini-2.5-flash")
    prompt = f"""Write exactly 3–5 sentences interpreting a company experience-curve analysis.
Company: {company} ({ticker}); annual observations: {len(frame)}; period: {frame['Fiscal Year'].min()}–{frame['Fiscal Year'].max()}.
Method: {method}. Fitted power exponent b={b:.5f}; learning rate={learning*100:.3f}%; log-log R squared={r2:.4f}.
Grounded interpretation: {base_text}
Requirements: Explicitly evaluate fit (R² < 0.5 means no visible reliable curve, 0.5–0.8 mixed, >0.8 strong descriptive fit). Discuss scale implications conditionally, never assert causation. Never invent competitor benchmarks, source data, average selling prices, or industry learning-rate claims. Mention cumulative units start within this data window, not lifetime. Do not make numerical calculations or change any supplied figures. Plain text only."""
    try:
        response = requests.post(
            f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent",
            params={"key": api_key},
            json={"contents": [{"parts": [{"text": prompt}]}], "generationConfig": {"temperature": 0.2, "maxOutputTokens": 650}},
            timeout=45,
        )
        response.raise_for_status()
        content = response.json()["candidates"][0]["content"]["parts"]
        text = " ".join(part.get("text", "") for part in content).strip()
        if not text:
            raise ValueError("Empty model response")
        return text, None
    except (requests.RequestException, KeyError, IndexError, ValueError) as err:
        return None, f"Gemini was unavailable ({err}); using the verified rule-based interpretation."

def pdf_report(company, ticker, frame, b, learning, r2, fitted, method, interpretation, cik, extra_note):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=(612, 792), rightMargin=42, leftMargin=42, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitleCustom", parent=styles["Heading1"], fontSize=17, leading=22, spaceAfter=10))
    styles.add(ParagraphStyle(name="ReportSmallCustom", parent=styles["BodyText"], fontSize=8.3, leading=11.5, spaceAfter=5))
    story = [Paragraph("Learning-Curve Analysis", styles["ReportTitleCustom"]),
             Paragraph(f"{company} ({ticker}) | SEC EDGAR | Generated {date.today().isoformat()}", styles["Normal"]), Spacer(1, 12)]
    def section(name):
        story.append(Paragraph(name, styles["Heading2"]))
    def paragraph(value):
        from xml.sax.saxutils import escape
        story.append(Paragraph(escape(str(value)), styles["ReportSmallCustom"]))
    section("Results")
    for label, value in [("Learning rate per doubling", f"{learning*100:.2f}%"), ("Coefficient b", f"{b:.5f}"), ("Log-log R squared", f"{r2:.4f}"), ("Annual observations", str(len(frame))), ("Fiscal years", f"{frame['Fiscal Year'].min()}–{frame['Fiscal Year'].max()}")]:
        paragraph(f"{label}: {value}")
    section("Log-log fitted curve")
    chart = png_chart(frame, fitted)
    story.append(Image(chart, width=7.1*inch, height=3.5*inch))
    section("Interpretation")
    paragraph(interpretation)
    section("Method and limitations")
    paragraph(f"Unit method: {method}. For the revenue proxy, annual units = reported revenue / modeled ASP, not directly observed physical sales. Cumulative units are summed from the first year of this window, not lifetime production. Unit cost = COGS / annual units. Fit ln(unit cost) = ln(A) - b × ln(cumulative units); learning rate = 1 - 2^(-b). R² is calculated in log space.")
    paragraph("This observational curve cannot isolate learning from inflation, product mix, acquisition effects, pricing, outsourcing, accounting changes, or capacity utilization. The historical window and chosen volume proxy materially affect the estimate.")
    if extra_note:
        paragraph(extra_note)
    section("Annual data and source filings")
    headers = ["FY", "Revenue ($M)", "COGS ($M)", "Annual units (M)", "Cum. units (M)", "ASP ($)", "Unit cost ($)"]
    table_rows = [headers]
    for _, r in frame.iterrows():
        table_rows.append([str(int(r["Fiscal Year"])), f'{r["Revenue ($)"]/1e6:,.1f}', f'{r["COGS ($)"]/1e6:,.1f}', f'{r["Annual Units"]/1e6:,.2f}', f'{r["Cumulative Units"]/1e6:,.2f}', f'{r["ASP ($)"]:.2f}' if pd.notna(r["ASP ($)"]) else "Direct", f'{r["Unit Cost ($)"]:.3f}'])
    t = Table(table_rows, colWidths=[31, 79, 75, 83, 83, 67, 93], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#e6eef5")), ("GRID", (0,0), (-1,-1), 0.25, colors.grey), ("FONTSIZE", (0,0), (-1,-1), 6.5), ("BOTTOMPADDING", (0,0), (-1,-1), 5), ("TOPPADDING", (0,0), (-1,-1), 5)]))
    story.append(t)
    paragraph("Annual and cumulative units are shown in millions; revenue and COGS in USD millions.")
    section("Unit-volume source and assumptions")
    # Consolidate identical sources; retain year-specific sources when they differ.
    source_groups = {}
    for _, r in frame.iterrows():
        source = str(r["Unit Source"]).strip()
        source_groups.setdefault(source, []).append(str(int(r["Fiscal Year"])))
    for source, years in source_groups.items():
        year_label = f"FY {years[0]}–{years[-1]}" if len(years) > 1 and len(years) == len(frame) else "FY " + ", ".join(years)
        paragraph(f"{year_label}: {source}")
    paragraph("Estimated ASPs are modeling inputs, not observed company-wide transaction prices. Catalog prices may differ from realized worldwide prices, and product mix and non-core business revenue can reduce the accuracy of inferred unit volumes.")
    section("SEC data and exact 10-K filings")
    paragraph("SEC Company Facts endpoint: https://data.sec.gov/api/xbrl/companyfacts/CIK" + f"{cik:010d}.json")
    for _, r in frame.iterrows():
        rev_link = accession_link(cik, r['Revenue Accession'])
        cost_link = accession_link(cik, r['COGS Accession'])
        if rev_link == cost_link:
            links = f"SEC filing: {rev_link}"
        else:
            links = f"Revenue filing: {rev_link}; COGS filing: {cost_link}"
        paragraph(f"FY {int(r['Fiscal Year'])}: Revenue 10-K {r['Revenue Filing']} ({r['Revenue Tag']}); COGS 10-K {r['COGS Filing']} ({r['COGS Tag']}). {links}")
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def word_report(company, ticker, frame, b, learning, r2, fitted, method, interpretation, cik, extra_note):
    """Create an editable DOCX report from the same regression results."""
    document = Document()
    sec = document.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(5)

    document.add_heading("Learning-Curve Analysis", 0)
    document.add_paragraph(f"{company} ({ticker}) | SEC EDGAR | Generated {date.today().isoformat()}")
    document.add_heading("Results", level=1)
    for label, value in [
        ("Learning rate per doubling", f"{learning*100:.2f}%"),
        ("Coefficient b", f"{b:.5f}"),
        ("Log-log R squared", f"{r2:.4f}"),
        ("Annual observations", str(len(frame))),
        ("Fiscal years", f"{frame['Fiscal Year'].min()}–{frame['Fiscal Year'].max()}"),
    ]:
        document.add_paragraph(f"{label}: {value}")

    document.add_heading("Log-log fitted curve", level=1)
    chart = png_chart(frame, fitted)
    document.add_picture(chart, width=Inches(6.7))
    document.add_heading("Interpretation", level=1)
    document.add_paragraph(interpretation)
    document.add_heading("Method and limitations", level=1)
    document.add_paragraph(
        f"Unit method: {method}. For the revenue proxy, annual units = reported revenue / modeled ASP, "
        "not directly observed physical sales. Cumulative units are summed from the first year "
        "of this window, not lifetime production. Unit cost = COGS / annual units. "
        "Fit ln(unit cost) = ln(A) - b × ln(cumulative units); learning rate = 1 - 2^(-b). "
        "R² is calculated in log space."
    )
    document.add_paragraph(
        "This observational curve cannot isolate learning from inflation, product mix, acquisition effects, "
        "pricing, outsourcing, accounting changes, or capacity utilization. The historical window "
        "and chosen volume proxy materially affect the estimate."
    )
    if extra_note:
        document.add_paragraph(extra_note)

    document.add_heading("Annual data and source filings", level=1)
    headers = ["FY", "Revenue ($M)", "COGS ($M)", "Annual units (M)", "Cum. units (M)", "ASP ($)", "Unit cost ($)"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for _, r in frame.iterrows():
        values = [str(int(r["Fiscal Year"])), f'{r["Revenue ($)"]/1e6:,.1f}',
                  f'{r["COGS ($)"]/1e6:,.1f}', f'{r["Annual Units"]/1e6:,.2f}',
                  f'{r["Cumulative Units"]/1e6:,.2f}',
                  f'{r["ASP ($)"]:.2f}' if pd.notna(r["ASP ($)"]) else "Direct",
                  f'{r["Unit Cost ($)"]:.3f}']
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(7)
    document.add_paragraph("Annual and cumulative units are shown in millions; revenue and COGS in USD millions.")

    document.add_heading("Unit-volume source and assumptions", level=1)
    source_groups = {}
    for _, r in frame.iterrows():
        source_groups.setdefault(str(r["Unit Source"]).strip(), []).append(str(int(r["Fiscal Year"])))
    for source, years in source_groups.items():
        year_label = f"FY {years[0]}–{years[-1]}" if len(years) == len(frame) else "FY " + ", ".join(years)
        document.add_paragraph(f"{year_label}: {source}")
    document.add_paragraph(
        "Estimated ASPs are modeling inputs, not observed company-wide transaction prices. "
        "Catalog prices may differ from realized worldwide prices, and product mix and non-core "
        "business revenue can reduce the accuracy of inferred unit volumes."
    )
    document.add_heading("SEC data and exact 10-K filings", level=1)
    document.add_paragraph(f"SEC Company Facts endpoint: https://data.sec.gov/api/xbrl/companyfacts/CIK{cik:010d}.json")
    for _, r in frame.iterrows():
        rev_link = accession_link(cik, r["Revenue Accession"])
        cost_link = accession_link(cik, r["COGS Accession"])
        links = f"SEC filing: {rev_link}" if rev_link == cost_link else f"Revenue filing: {rev_link}; COGS filing: {cost_link}"
        document.add_paragraph(
            f"FY {int(r['Fiscal Year'])}: Revenue 10-K {r['Revenue Filing']} ({r['Revenue Tag']}); "
            f"COGS 10-K {r['COGS Filing']} ({r['COGS Tag']}). {links}"
        )
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()

# ----- Inputs -----
left, right = st.columns([2, 1])
with left:
    ticker = st.text_input("Public company ticker", "PG").strip().upper()
with right:
    years = st.selectbox("Annual fiscal years", [5, 10, 12, 15], index=1)

st.info("The SEC API generally provides revenue and COGS, **not comparable units sold**. The tool will require annual unit counts or documented annual average selling prices after retrieving the financials.")

if st.button("1. Retrieve SEC financials", type="primary"):
    try:
        with st.spinner("Loading SEC company facts..."):
            company = find_company(ticker)
            if not company:
                st.error("Ticker not found in SEC ticker directory.")
                st.stop()
            cik = int(company["cik_str"])
            facts = get_facts(cik)
            cogs = annual_facts(facts, COGS_TAGS, years)
            revenue = annual_facts(facts, REV_TAGS, years)
            if cogs.empty or revenue.empty:
                st.error("No usable annual SEC COGS or revenue facts found for this ticker. This ticker may require custom tags or manual verification.")
                st.stop()
            merged = pd.merge(revenue, cogs, on="Period End", suffixes=("_Revenue", "_COGS"))
            merged = merged.sort_values("Period End").tail(years).reset_index(drop=True)
            if len(merged) < 5:
                st.error(f"Only {len(merged)} matching annual COGS/revenue observations were found. At least five are required.")
                st.stop()
            st.session_state["financials"] = merged
            st.session_state["company"] = company["title"]
            st.session_state["cik"] = cik
            st.session_state["ticker"] = ticker
            st.session_state.pop("results", None)
            st.session_state.pop("ai_interpretation", None)
    except requests.RequestException as err:
        st.error(f"SEC retrieval failed: {err}. Check your SEC contact email and network access.")
    except Exception as err:
        st.error(f"Could not process SEC data: {err}")

if "financials" in st.session_state:
    raw = st.session_state["financials"].copy()
    company = st.session_state["company"]
    cik = st.session_state["cik"]
    st.success(f"Loaded {len(raw)} matched fiscal years for {company} ({st.session_state['ticker']}).")
    st.warning("Check extracted financial values and XBRL tags against the linked 10-K filings before relying on the result. Earliest *available* filings may include comparative figures.")
    with st.expander("Review SEC financial observations and exact filings", expanded=False):
        st.dataframe(raw, use_container_width=True, hide_index=True)
        for _, row in raw.iterrows():
            st.markdown(f"FY {int(row['Fiscal Year_Revenue'])}: [Revenue filing]({accession_link(cik, row['Accession_Revenue'])}) · [COGS filing]({accession_link(cik, row['Accession_COGS'])})")

    st.subheader("2. Document unit volumes")
    method = st.radio("Unit-volume method", ["Revenue / estimated annual average selling prices", "Direct annual units sold"], horizontal=True)
    editable = pd.DataFrame({
        "Fiscal Year": raw["Fiscal Year_Revenue"].astype(int),
        "Revenue ($)": raw["Value_Revenue"],
        "COGS ($)": raw["Value_COGS"],
        "Average Selling Price ($)": [None] * len(raw),
        "Annual Units Sold": [None] * len(raw),
        "Unit Source / URL": [""] * len(raw),
    })
    if method.startswith("Revenue"):
        st.caption("Provide a defensible **company-wide estimated** average selling price for EACH year, with its source and assumptions. A single product price is not a company-wide average.")
        disabled = ["Fiscal Year", "Revenue ($)", "COGS ($)", "Annual Units Sold"]
    else:
        st.caption("Enter comparable annual unit counts with a source for EACH year. The unit definition must be consistent over time.")
        disabled = ["Fiscal Year", "Revenue ($)", "COGS ($)", "Average Selling Price ($)"]
    edited = st.data_editor(editable, hide_index=True, num_rows="fixed", disabled=disabled,
                            column_config={"Average Selling Price ($)": st.column_config.NumberColumn(min_value=0.000001, format="%.4f"),
                                           "Annual Units Sold": st.column_config.NumberColumn(min_value=0.000001, format="%.0f"),
                                           "Unit Source / URL": st.column_config.TextColumn(width="large")},
                            key=f"editor_{st.session_state['ticker']}_{len(raw)}_{method}", use_container_width=True)
    note = st.text_area("Optional price-guide URL / additional comparability notes (printed once in Word report)", placeholder="Paste the official price-guide URL and any additional assumptions or filing differences.")
    if st.button("3. Calculate curve and generate report", type="primary"):
        try:
            work = edited.copy()
            if (work["Unit Source / URL"].fillna("").astype(str).str.strip() == "").any():
                raise ValueError("Provide a unit-volume or ASP source for every year.")
            if method.startswith("Revenue"):
                prices = pd.to_numeric(work["Average Selling Price ($)"], errors="coerce")
                if prices.isna().any() or (prices <= 0).any():
                    raise ValueError("Enter a positive estimated average selling price for every year.")
                work["ASP ($)"] = prices
                work["Annual Units"] = work["Revenue ($)"] / prices
            else:
                units = pd.to_numeric(work["Annual Units Sold"], errors="coerce")
                if units.isna().any() or (units <= 0).any():
                    raise ValueError("Enter positive annual unit counts for every year.")
                work["Annual Units"] = units
                work["ASP ($)"] = np.nan
            work["Unit Source"] = work["Unit Source / URL"]
            work["Cumulative Units"] = work["Annual Units"].cumsum()
            work["Unit Cost ($)"] = work["COGS ($)"] / work["Annual Units"]
            work["Revenue Filing"] = raw["Accession_Revenue"]
            work["COGS Filing"] = raw["Accession_COGS"]
            work["Revenue Accession"] = raw["Accession_Revenue"]
            work["COGS Accession"] = raw["Accession_COGS"]
            work["Revenue Tag"] = raw["XBRL Tag_Revenue"]
            work["COGS Tag"] = raw["XBRL Tag_COGS"]
            if not np.isfinite(work["Unit Cost ($)"].to_numpy(dtype=float)).all():
                raise ValueError("Non-finite unit costs; check inputs.")
            b, learning, r2, fitted, A = regression(work)
            text = assessment(b, learning, r2, method)
            st.session_state.pop("ai_interpretation", None)
            st.session_state["results"] = (work, b, learning, r2, fitted, A, method, text, note)
        except (ValueError, TypeError, OverflowError, np.linalg.LinAlgError) as err:
            st.error(str(err))

if "results" in st.session_state:
    work, b, learning, r2, fitted, A, method, rule_text, note = st.session_state["results"]
    st.divider()
    st.subheader("Learning-curve results")
    c1, c2, c3 = st.columns(3)
    c1.metric("Learning rate / doubling", f"{learning*100:.2f}%")
    c2.metric("Coefficient b", f"{b:.5f}")
    c3.metric("Log-log R²", f"{r2:.4f}")
    if r2 < 0.5:
        st.error("Weak fit: no reliable experience curve is visible in this selected window.")
    elif r2 < 0.8:
        st.warning("Mixed fit: interpret the estimated learning rate cautiously.")
    else:
        st.success("Strong descriptive fit. This does not establish that experience caused the cost changes.")
    st.plotly_chart(plotly_chart(work, fitted), use_container_width=True)
    st.latex(r"\mathrm{Cost}(N)=A\,N^{-b}\qquad\mathrm{Learning\ rate}=1-2^{-b}")
    st.caption(f"Fitted A = {A:.6g}; N is cumulative units measured from the FIRST year in this data window, not lifetime production. R² is computed from logged values.")
    st.subheader("Annual calculation table")
    st.dataframe(work[["Fiscal Year", "Revenue ($)", "COGS ($)", "ASP ($)", "Annual Units", "Cumulative Units", "Unit Cost ($)", "Unit Source"]], hide_index=True, use_container_width=True)
    st.subheader("Strategic interpretation")
    st.write(rule_text)
    ai_text = None
    if st.button("Generate optional Gemini interpretation"):
        with st.spinner("Requesting Gemini analysis..."):
            ai_text, error = gemini_analysis(st.session_state["company"], st.session_state["ticker"], work, b, learning, r2, method, rule_text)
            if error:
                st.warning(error)
            if ai_text:
                st.session_state["ai_interpretation"] = ai_text
    if st.session_state.get("ai_interpretation"):
        st.markdown("**Gemini-generated interpretation (review for accuracy):**")
        st.write(st.session_state["ai_interpretation"])
    chosen_text = st.session_state.get("ai_interpretation") or rule_text
    st.subheader("Download report")
    try:
        docx_bytes = word_report(st.session_state["company"], st.session_state["ticker"], work, b, learning, r2, fitted, method, chosen_text, st.session_state["cik"], note)
        st.download_button("📄 Download learning-curve Word document", data=docx_bytes,
                           file_name=f"{st.session_state['ticker']}_learning_curve_report.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except Exception as err:
        st.error(f"Word document generation failed: {err}")
    st.caption("For your assignment, place the public app URL at the top of the final submission PDF and append your own 250–500-word reflection.")
